import csv
import io
import uuid
from rest_framework import viewsets, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import AllowAny, IsAuthenticated, SAFE_METHODS
from django.contrib.auth import update_session_auth_hash
from rest_framework.decorators import action
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework_simplejwt.serializers import TokenObtainPairSerializer
from rest_framework_simplejwt.views import TokenObtainPairView
from .models import (Institution, User, Course, Enrollment, Grade, ClassSession,
                     AttendanceRecord, ProjectGroup, Task, Quiz, Question, QuizAttempt, AttemptAnswer,
                     Assignment, Submission, Material)
from .serializers import (InstitutionSerializer, UserSerializer, CourseSerializer,
                          EnrollmentSerializer, GradeSerializer, ClassSessionSerializer,
                          AttendanceRecordSerializer, ProjectGroupSerializer, TaskSerializer,
                           QuestionSerializer, QuizSerializer, QuizAttemptSerializer, AttemptAnswerSerializer,
                           AssignmentSerializer, SubmissionSerializer, MaterialSerializer)
from .permissions import IsAdmin, IsAdminOrInstructor, IsInstructor, IsStudent


class StudentDashboardView(APIView):
    permission_classes = [IsStudent]

    def get(self, request):
        enrollments = Enrollment.objects.filter(student=request.user).select_related('course')
        courses = []
        total_gpa = 0.0
        grade_count = 0
        now = timezone.now()

        for e in enrollments:
            c = e.course
            grades = Grade.objects.filter(student=request.user, course=c)
            course_grades = [g.grade for g in grades]
            numeric = []
            for g in course_grades:
                try:
                    numeric.append(float(g))
                except (ValueError, TypeError):
                    pass
            avg = sum(numeric) / len(numeric) if numeric else 0
            if avg:
                total_gpa += avg
                grade_count += 1

            sessions = ClassSession.objects.filter(course=c)
            attendance_records = AttendanceRecord.objects.filter(student=request.user, session__course=c)
            total_sessions = sessions.count()
            attended = attendance_records.filter(status__in=['present', 'late']).count()

            courses.append({
                'id': c.id,
                'title': c.title,
                'course_code': c.course_code,
                'instructor_name': c.instructor.get_full_name() or c.instructor.username if c.instructor else 'N/A',
                'average_grade': round(avg, 1) if avg else None,
                'attendance_pct': round((attended / total_sessions) * 100) if total_sessions else None,
            })

        gpa = round(total_gpa / grade_count, 2) if grade_count else 0.0

        upcoming_sessions = ClassSession.objects.filter(
            course__enrolled_students__student=request.user,
            scheduled_at__gte=now,
        ).order_by('scheduled_at')[:3]

        due_assignments = Assignment.objects.filter(
            course__enrolled_students__student=request.user,
        ).order_by('due_date')[:7]

        next_session = upcoming_sessions.first()

        return Response({
            'courses': courses,
            'gpa': gpa,
            'enrolled_count': len(courses),
            'pending_tasks': due_assignments.count(),
            'next_session': {
                'id': next_session.id,
                'title': next_session.title,
                'scheduled_at': next_session.scheduled_at,
                'meeting_link': next_session.meeting_link,
                'course': next_session.course.title,
            } if next_session else None,
            'due_assignments': [{
                'id': a.id,
                'title': a.title,
                'course': a.course.title,
                'course_id': a.course.id,
                'due_date': a.due_date,
                'max_points': a.max_points,
                'assignment_type': a.assignment_type,
            } for a in due_assignments],
        })


class ProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        serializer = UserSerializer(request.user)
        return Response(serializer.data)

    def patch(self, request):
        user = request.user
        allowed = ['first_name', 'last_name', 'email']
        data = {k: v for k, v in request.data.items() if k in allowed}
        serializer = UserSerializer(user, data=data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def post(self, request):
        user = request.user
        old = request.data.get('old_password')
        new = request.data.get('new_password')
        if not old or not new:
            return Response({'detail': 'Both old and new password required.'}, status=status.HTTP_400_BAD_REQUEST)
        if not user.check_password(old):
            return Response({'detail': 'Old password is incorrect.'}, status=status.HTTP_400_BAD_REQUEST)
        user.set_password(new)
        user.save()
        update_session_auth_hash(request, user)
        return Response({'detail': 'Password changed successfully.'})


class BulkUserImportView(APIView):
    permission_classes = [IsAdmin]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'No file provided.'}, status=status.HTTP_400_BAD_REQUEST)

        results = {'created': 0, 'skipped': 0, 'errors': []}
        default_password = 'Temp@1234'

        try:
            if file.name.endswith('.csv'):
                decoded = file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(decoded))
                rows = list(reader)
            else:
                from openpyxl import load_workbook
                wb = load_workbook(file, read_only=True)
                ws = wb.active
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    rows.append(dict(zip(headers, row)))
        except Exception as e:
            return Response({'detail': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        for i, row in enumerate(rows, start=2):
            first_name = (row.get('first_name') or '').strip()
            last_name = (row.get('last_name') or '').strip()
            email = (row.get('email') or '').strip()
            student_id = (row.get('student_id') or row.get('username') or '').strip()

            if not first_name or not last_name or not email or not student_id:
                results['errors'].append(f"Row {i}: missing required fields")
                results['skipped'] += 1
                continue

            if User.objects.filter(username=student_id).exists():
                results['skipped'] += 1
                continue

            try:
                User.objects.create_user(
                    username=student_id,
                    email=email,
                    first_name=first_name,
                    last_name=last_name,
                    password=default_password,
                    role='student',
                    institution=request.user.institution
                )
                results['created'] += 1
            except Exception as e:
                results['errors'].append(f"Row {i}: {str(e)}")
                results['skipped'] += 1

        return Response(results)


class DownloadTemplateView(APIView):
    permission_classes = [IsAdmin]

    def get(self, request):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="bulk_import_template.csv"'
        writer = csv.writer(response)
        writer.writerow(['first_name', 'last_name', 'email', 'student_id'])
        writer.writerow(['John', 'Doe', 'john.doe@university.edu', 'STU001'])
        writer.writerow(['Jane', 'Smith', 'jane.smith@university.edu', 'STU002'])
        return response


class MyInstitutionSettingsView(APIView):
    permission_classes = [IsAdmin]

    def get(self, request):
        institution = request.user.institution
        if not institution:
            return Response({'detail': 'No institution assigned.'}, status=status.HTTP_404_NOT_FOUND)
        serializer = InstitutionSerializer(institution)
        return Response(serializer.data)

    def patch(self, request):
        institution = request.user.institution
        if not institution:
            return Response({'detail': 'No institution assigned.'}, status=status.HTTP_404_NOT_FOUND)
        serializer = InstitutionSerializer(institution, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class InstitutionViewSet(viewsets.ModelViewSet):
    queryset = Institution.objects.all()
    serializer_class = InstitutionSerializer

    def get_permissions(self):
        if self.action in ['create', 'update', 'partial_update', 'destroy']:
            return [IsAdmin()]
        return [IsAuthenticated()]


class UserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAdmin]

    def get_queryset(self):
        user = self.request.user
        qs = User.objects.all()
        if user.role != 'super_admin' and user.institution:
            qs = qs.filter(institution=user.institution)
        elif user.role != 'super_admin':
            return User.objects.none()
        search = self.request.query_params.get('search', None)
        if search:
            qs = qs.filter(
                Q(username__icontains=search) |
                Q(email__icontains=search) |
                Q(first_name__icontains=search) |
                Q(last_name__icontains=search)
            )
        return qs

    def perform_create(self, serializer):
        serializer.save(institution=self.request.user.institution)


class CourseViewSet(viewsets.ModelViewSet):
    queryset = Course.objects.all()
    serializer_class = CourseSerializer

    def get_permissions(self):
        if self.request.method in SAFE_METHODS:
            return [IsAuthenticated()]
        if self.action in ['enroll', 'generate_key']:
            return [IsAuthenticated()]
        return [IsAdminOrInstructor()]

    def get_queryset(self):
        user = self.request.user
        qs = Course.objects.all()
        if user.role == 'student':
            return qs.filter(institution=user.institution)
        if user.role != 'super_admin' and user.institution:
            qs = qs.filter(institution=user.institution)
        elif user.role != 'super_admin':
            return Course.objects.none()
        search = self.request.query_params.get('search', None)
        if search:
            qs = qs.filter(
                Q(title__icontains=search) |
                Q(course_code__icontains=search) |
                Q(description__icontains=search)
            )
        return qs

    def perform_create(self, serializer):
        serializer.save(institution=self.request.user.institution)

    @action(detail=True, methods=['post'])
    def enroll(self, request, pk=None):
        course = self.get_object()
        key = request.data.get('enrollment_key', '')
        if course.enrollment_key and course.enrollment_key != key:
            return Response({'detail': 'Invalid enrollment key.'}, status=status.HTTP_400_BAD_REQUEST)
        if Enrollment.objects.filter(student=request.user, course=course).exists():
            return Response({'detail': 'Already enrolled.'}, status=status.HTTP_400_BAD_REQUEST)
        Enrollment.objects.create(student=request.user, course=course)
        return Response({'detail': 'Enrolled successfully.'})

    @action(detail=True, methods=['post'])
    def generate_key(self, request, pk=None):
        course = self.get_object()
        course.enrollment_key = uuid.uuid4().hex[:8].upper()
        course.save(update_fields=['enrollment_key'])
        return Response({'enrollment_key': course.enrollment_key})


class MaterialViewSet(viewsets.ModelViewSet):
    queryset = Material.objects.all()
    serializer_class = MaterialSerializer

    def get_permissions(self):
        if self.request.method in SAFE_METHODS:
            return [IsAuthenticated()]
        return [IsAdminOrInstructor()]

    def get_queryset(self):
        user = self.request.user
        qs = Material.objects.all()
        if user.role == 'student':
            return qs.filter(course__enrolled_students__student=user)
        if user.role == 'super_admin':
            return qs
        if user.role == 'instructor':
            return qs.filter(course__instructor=user)
        if user.institution:
            return qs.filter(course__institution=user.institution)
        return Material.objects.none()

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class InstructorDashboardView(APIView):
    permission_classes = [IsInstructor]

    def get(self, request):
        courses = Course.objects.filter(instructor=request.user)
        course_data = []
        for c in courses:
            enrolled = Enrollment.objects.filter(course=c).count()
            course_data.append({
                'id': c.id,
                'title': c.title,
                'course_code': c.course_code,
                'students': enrolled,
            })
        return Response({
            'courses': course_data,
        })


class DashboardStatsView(APIView):
    permission_classes = [IsAdmin]

    def get(self, request):
        users_count = User.objects.count()
        courses_count = Course.objects.count()
        enrollments_count = Enrollment.objects.count()
        recent_courses = Course.objects.order_by('-created_at')[:5]
        recent_users = User.objects.order_by('-date_joined')[:5]
        activities = []
        for c in recent_courses:
            activities.append({
                'id': f'course-{c.id}',
                'action': f'Created course: {c.title} ({c.course_code})',
                'user': c.instructor.username if c.instructor else 'System',
                'time': c.created_at.isoformat() if c.created_at else '',
            })
        for u in recent_users:
            activities.append({
                'id': f'user-{u.id}',
                'action': f'Registered user: {u.username} ({u.get_role_display()})',
                'user': 'System',
                'time': u.date_joined.isoformat() if u.date_joined else '',
            })
        activities.sort(key=lambda x: x['time'], reverse=True)
        return Response({
            'total_users': users_count,
            'total_courses': courses_count,
            'total_enrollments': enrollments_count,
            'recent_activities': activities[:10],
        })


class CustomTokenObtainPairSerializer(TokenObtainPairSerializer):
    def validate(self, attrs):
        data = super().validate(attrs)
        data['role'] = self.user.role
        data['user_id'] = self.user.id
        data['username'] = self.user.username
        return data


class CustomTokenObtainPairView(TokenObtainPairView):
    serializer_class = CustomTokenObtainPairSerializer


class EnrollmentViewSet(viewsets.ModelViewSet):
    queryset = Enrollment.objects.all()
    serializer_class = EnrollmentSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return Enrollment.objects.all()
        if user.institution:
            return Enrollment.objects.filter(student__institution=user.institution)
        return Enrollment.objects.none()


class GradeViewSet(viewsets.ModelViewSet):
    queryset = Grade.objects.all()
    serializer_class = GradeSerializer

    def get_queryset(self):
        user = self.request.user
        qs = Grade.objects.all()
        if user.role == 'super_admin':
            pass
        elif user.institution:
            qs = qs.filter(student__institution=user.institution)
        else:
            return Grade.objects.none()
        course = self.request.query_params.get('course')
        if course:
            qs = qs.filter(course_id=course)
        return qs


class ClassSessionViewSet(viewsets.ModelViewSet):
    queryset = ClassSession.objects.all()
    serializer_class = ClassSessionSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return ClassSession.objects.all()
        if user.institution:
            return ClassSession.objects.filter(course__institution=user.institution)
        return ClassSession.objects.none()

    def perform_create(self, serializer):
        serializer.save()

    @action(detail=True, methods=['post'])
    def mark_all_present(self, request, pk=None):
        session = self.get_object()
        enrollments = Enrollment.objects.filter(course=session.course)
        created = 0
        for enrollment in enrollments:
            _, was_created = AttendanceRecord.objects.get_or_create(
                session=session,
                student=enrollment.student,
                defaults={'status': 'present', 'is_verified': True}
            )
            if was_created:
                created += 1
            else:
                AttendanceRecord.objects.filter(session=session, student=enrollment.student).update(
                    status='present', is_verified=True
                )
        return Response({'marked_present': enrollments.count(), 'created': created})

    @action(detail=True, methods=['get'])
    def stats(self, request, pk=None):
        session = self.get_object()
        total = AttendanceRecord.objects.filter(session=session).count()
        present = AttendanceRecord.objects.filter(session=session, status='present').count()
        absent = AttendanceRecord.objects.filter(session=session, status='absent').count()
        late = AttendanceRecord.objects.filter(session=session, status='late').count()
        excused = AttendanceRecord.objects.filter(session=session, status='excused').count()
        return Response({
            'total': total,
            'present': present,
            'absent': absent,
            'late': late,
            'excused': excused,
        })


class AttendanceRecordViewSet(viewsets.ModelViewSet):
    queryset = AttendanceRecord.objects.all()
    serializer_class = AttendanceRecordSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return AttendanceRecord.objects.all()
        if user.institution:
            return AttendanceRecord.objects.filter(student__institution=user.institution)
        return AttendanceRecord.objects.none()

    def get_serializer(self, *args, **kwargs):
        if isinstance(kwargs.get('data'), list):
            kwargs['many'] = True
        return super().get_serializer(*args, **kwargs)


class ProjectGroupViewSet(viewsets.ModelViewSet):
    queryset = ProjectGroup.objects.all()
    serializer_class = ProjectGroupSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return ProjectGroup.objects.all()
        if user.institution:
            return ProjectGroup.objects.filter(course__institution=user.institution)
        return ProjectGroup.objects.none()

    def perform_create(self, serializer):
        serializer.save()

    @action(detail=False, methods=['post'])
    def random_groups(self, request):
        course_id = request.data.get('course')
        n_groups = int(request.data.get('n_groups', 2))
        if not course_id:
            return Response({'detail': 'course is required'}, status=status.HTTP_400_BAD_REQUEST)

        enrollments = Enrollment.objects.filter(course_id=course_id).select_related('student')
        students = [e.student for e in enrollments]

        import random
        random.shuffle(students)

        groups = []
        for i in range(n_groups):
            group_name = f"Group {chr(65 + i)}"
            group, created = ProjectGroup.objects.get_or_create(
                course_id=course_id,
                name=f"Group {chr(65 + i)}",
            )
            groups.append(group)

        # Clear existing members and reassign
        for g in groups:
            g.members.clear()

        for idx, student in enumerate(students):
            groups[idx % n_groups].members.add(student)

        serializer = ProjectGroupSerializer(groups, many=True)
        return Response(serializer.data)


class TaskViewSet(viewsets.ModelViewSet):
    queryset = Task.objects.all()
    serializer_class = TaskSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return Task.objects.all()
        if user.institution:
            return Task.objects.filter(group__course__institution=user.institution)
        return Task.objects.none()


class QuestionViewSet(viewsets.ModelViewSet):
    queryset = Question.objects.all()
    serializer_class = QuestionSerializer
    permission_classes = [IsAdminOrInstructor]

    def get_queryset(self):
        user = self.request.user
        qs = Question.objects.all()
        if user.role == 'super_admin':
            return qs
        if user.institution:
            return qs.filter(quiz__course__institution=user.institution)
        return Question.objects.none()


class QuizViewSet(viewsets.ModelViewSet):
    queryset = Quiz.objects.all()
    serializer_class = QuizSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return Quiz.objects.all()
        if user.institution:
            return Quiz.objects.filter(course__institution=user.institution)
        return Quiz.objects.none()

    def perform_create(self, serializer):
        serializer.save(instructor=self.request.user)


class QuizAttemptViewSet(viewsets.ModelViewSet):
    queryset = QuizAttempt.objects.all()
    serializer_class = QuizAttemptSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return QuizAttempt.objects.all()
        if user.institution:
            return QuizAttempt.objects.filter(student__institution=user.institution)
        return QuizAttempt.objects.none()


class AttemptAnswerViewSet(viewsets.ModelViewSet):
    queryset = AttemptAnswer.objects.all()
    serializer_class = AttemptAnswerSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'super_admin':
            return AttemptAnswer.objects.all()
        if user.institution:
            return AttemptAnswer.objects.filter(attempt__student__institution=user.institution)
        return AttemptAnswer.objects.none()


class AssignmentViewSet(viewsets.ModelViewSet):
    queryset = Assignment.objects.all()
    serializer_class = AssignmentSerializer

    def get_queryset(self):
        user = self.request.user
        qs = Assignment.objects.all()
        if user.role == 'super_admin':
            pass
        elif user.institution:
            qs = qs.filter(course__institution=user.institution)
        else:
            return Assignment.objects.none()
        course = self.request.query_params.get('course')
        if course:
            qs = qs.filter(course_id=course)
        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


class SubmissionViewSet(viewsets.ModelViewSet):
    queryset = Submission.objects.all()
    serializer_class = SubmissionSerializer

    def get_queryset(self):
        user = self.request.user
        qs = Submission.objects.all()
        if user.role == 'super_admin':
            pass
        elif user.institution:
            qs = qs.filter(assignment__course__institution=user.institution)
        else:
            return Submission.objects.none()
        assignment = self.request.query_params.get('assignment')
        if assignment:
            qs = qs.filter(assignment_id=assignment)
        return qs
